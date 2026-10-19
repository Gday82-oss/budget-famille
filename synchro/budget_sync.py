#!/usr/bin/env python3
"""Synchronisation dans les deux sens entre le classeur Budget_Guillaume-Anne.xlsx (PC)
et l'appli Budget (Supabase, téléphones).

  python3 budget_sync.py connexion   # première fois : e-mail, mot de passe, chemin du fichier
  python3 budget_sync.py essai       # montre ce qui serait synchronisé, sans rien modifier
  python3 budget_sync.py une-fois    # fait une synchronisation puis s'arrête
  python3 budget_sync.py             # surveille en continu (lancé par le service au démarrage)

Règles :
- Chaque mois = une feuille « Budget mensuel <mois> <année> ». Les lignes sont reconnues par leur nom.
- Le « Coût réel » est une liste de montants (=a+b+c) : chaque montant est une opération dans l'appli.
- On compare l'Excel et l'appli à la dernière version synchronisée (fusion à trois voies) :
  un ajout ou une suppression d'un côté est recopié de l'autre, rien n'est perdu.
- Si la même valeur (montant prévu, commentaire) a changé des deux côtés, l'Excel l'emporte.
- Si le fichier est ouvert (LibreOffice), les changements de l'appli attendent sa fermeture.
- Une sauvegarde du classeur est faite avant chaque écriture.
"""
import argparse
import getpass
import json
import logging
import os
import re
import shutil
import subprocess
import sys
import tempfile
import time
import urllib.error
import urllib.request
import warnings
from collections import Counter
from datetime import date, datetime
from pathlib import Path

try:
    import openpyxl
    from openpyxl.workbook.properties import CalcProperties
except ImportError:  # pragma: no cover
    sys.exit("Il manque openpyxl : sudo apt install python3-openpyxl python3-pil")

SUPABASE_URL = os.environ.get("BUDGET_SUPABASE_URL", "https://sahxtcsblspjzhipuqyq.supabase.co")
SUPABASE_KEY = os.environ.get("BUDGET_SUPABASE_KEY", "sb_publishable_AhIuVcOzoCudCzALjWBdLQ_IeU_qmef")
DOSSIER_CONFIG = Path(os.environ.get("BUDGET_SYNC_CONFIG", Path.home() / ".config" / "budget-sync"))
DOSSIER_DONNEES = Path(os.environ.get("BUDGET_SYNC_DATA", Path.home() / ".local" / "share" / "budget-sync"))
FICHIER_DEFAUT = Path.home() / "Bureau" / "Budget_Guillaume-Anne.xlsx"
INTERVALLE_FICHIER = 3     # secondes entre deux regards sur le fichier
INTERVALLE_APPLI = 10      # secondes entre deux regards sur l'appli
SAUVEGARDES_GARDEES = 40

log = logging.getLogger("budget-sync")
warnings.filterwarnings("ignore", module="openpyxl")

MOIS_NOMS = ["janvier", "février", "mars", "avril", "mai", "juin", "juillet", "août",
             "septembre", "octobre", "novembre", "décembre"]
MOIS_LUS = {n: i + 1 for i, n in enumerate(MOIS_NOMS)}
MOIS_LUS.update({"fevrier": 2, "aout": 8, "decembre": 12, "sept": 9, "janv": 1, "fev": 2, "févr": 2,
                 "avr": 4, "juil": 7, "oct": 10, "nov": 11, "dec": 12, "déc": 12})
SECTIONS = {"DÉPENSES FIXES": "fixe", "DEPENSES FIXES": "fixe", "DÉPENSES VARIABLES": "variable",
            "DEPENSES VARIABLES": "variable", "REVENUS": "revenu"}
GROUPES = {"ENFANTS": "Enfants", "FRAIS LOGEMENT + TÉLÉPHONE": "Logement & téléphone",
           "ASSURANCES – IMPÔTS": "Assurances & impôts", "CRÉDITS": "Crédits", "AUTRES": "Autres",
           "COURSES ALIMENTATION/ LOGEMENT": "Courses", "VOITURES": "Voitures", "LOISIRS ET DIVERS": "Loisirs & divers"}
LIBELLE_REPORT = "Solde réel avec reprise solde mois précédent"


# ---------------------------------------------------------------- outils
def normaliser(texte):
    return " ".join(str(texte).split()).casefold()


def centimes(x):
    return int(round(float(x) * 100))


def en_euros(c):
    return round(c / 100, 2)


def texte_montant(c):
    s = f"{abs(c) / 100:.2f}".rstrip("0").rstrip(".")
    return s


def notifier(message):
    log.warning(message)
    try:
        subprocess.run(["notify-send", "-a", "Budget", "Budget – synchronisation", message],
                       check=False, timeout=5, capture_output=True)
    except Exception:
        pass


def lire_montants(v):
    """Cellule « Coût réel » -> liste de centimes ; None si la formule n'est pas une simple somme."""
    if v is None or (isinstance(v, str) and not v.strip()):
        return []
    if isinstance(v, (int, float)):
        return [] if v == 0 else [centimes(v)]
    s = str(v).strip().lstrip("=").replace(" ", "").replace(",", ".")
    if not s:
        return []
    morceaux = re.findall(r"[+-]?[^+-]+", s)
    if not all(re.fullmatch(r"[+-]?\d+(?:\.\d+)?", m) for m in morceaux):
        return None
    return [c for c in (centimes(m) for m in morceaux) if c != 0]


def lire_prevu(v):
    """Cellule « Coût prévu » -> euros, None si vide, '?' si formule non simple."""
    if v is None or (isinstance(v, str) and not v.strip()):
        return None
    m = lire_montants(v)
    if m is None:
        return "?"
    return en_euros(sum(m))


def lire_taux(v):
    if v is None or (isinstance(v, str) and not v.strip()):
        return 0.0
    s = str(v).replace(",", ".")
    m = re.search(r"\*\s*(\d+(?:\.\d+)?)\s*%", s)
    if m:
        return round(float(m.group(1)) / 100, 6)
    m = re.search(r"\*\s*(0?\.\d+)", s)
    if m:
        return round(float(m.group(1)), 6)
    return None


def lire_report(v):
    if v is None or (isinstance(v, str) and not v.strip()):
        return 0.0
    if isinstance(v, (int, float)):
        return round(float(v), 2)
    m = re.fullmatch(r"=\s*D\d+\s*([+-])\s*(\d+(?:[.,]\d+)?)\s*", str(v))
    if m:
        return round(float(m.group(2).replace(",", ".")) * (1 if m.group(1) == "+" else -1), 2)
    return None


def valeur_montants(liste):
    if not liste:
        return None
    if len(liste) == 1:
        return en_euros(liste[0])
    f = "="
    for i, c in enumerate(liste):
        f += ("-" if c < 0 else ("" if i == 0 else "+")) + texte_montant(c)
    return f


def mois_de_feuille(titre):
    m = re.fullmatch(r"\s*budget\s+mensuel\s+(\S+)\s+(\d{4})\s*", titre, re.I)
    if not m or m.group(1).casefold() not in MOIS_LUS:
        return None
    return f"{m.group(2)}-{MOIS_LUS[m.group(1).casefold()]:02d}"


def nom_feuille(mois_id):
    a, m = mois_id.split("-")
    return f"Budget mensuel {MOIS_NOMS[int(m) - 1]} {a}"


def nom_mois(mois_id):
    a, m = mois_id.split("-")
    return f"{MOIS_NOMS[int(m) - 1]} {a}"


def fusion_valeur(base, excel, appli):
    """Fusion d'une valeur simple : le côté qui a changé gagne ; si les deux ont changé, l'Excel gagne."""
    if excel == base:
        return appli
    return excel


# ---------------------------------------------------------------- lecture du classeur
class Feuille:
    def __init__(self, ws, mois_id):
        self.ws, self.mois_id = ws, mois_id
        self.lignes = {}          # clé -> dict(rang, libelle, bloc, groupe, categorie, prevu, montants, commentaire)
        self.rang_total = self.rang_solde = self.rang_report = self.rang_imprevus = None
        self.report, self.taux = None, None
        self._lire()

    def _lire(self):
        ws, bloc, groupe = self.ws, None, None
        for r in range(2, ws.max_row + 1):
            a = ws.cell(r, 1).value
            if a is None or not str(a).strip():
                continue
            libelle = " ".join(str(a).split())
            haut = libelle.upper()
            b, c, d, f = (ws.cell(r, k).value for k in (2, 3, 4, 6))
            if haut.startswith(("SOLDE RÉEL", "SOLDE REEL")) or "REPRISE" in haut:
                self.rang_report, self.report = r, lire_report(d)
                continue
            if re.match(r"SOUS[ -]TOTAL D[ÉE]PENSES TOTALES|D[ÉE]PENSES TOTALES", haut):
                self.rang_total, bloc, groupe = r, "provision", "Épargne & divers"
                continue
            if haut.startswith(("SOUS TOTAL", "SOUS-TOTAL", "TOTAL")):
                continue
            if haut.startswith("SOLDE"):
                self.rang_solde = r
                continue
            if haut in SECTIONS:
                bloc = SECTIONS[haut]
                groupe = "Revenus" if bloc == "revenu" else None
                continue
            imprevus = haut.startswith(("IMPRÉVUS", "IMPREVUS"))
            if not imprevus and bloc in ("fixe", "variable") and b is None and c is None and d is None:
                groupe = GROUPES.get(haut, libelle[:1].upper() + libelle[1:].lower())
                continue
            if bloc is None or groupe is None:
                continue
            if imprevus:
                self.rang_imprevus, self.taux = r, lire_taux(c)
            cle = "imprevus" if imprevus else normaliser(libelle)
            self.lignes[cle] = dict(
                rang=r, libelle="Imprévus" if imprevus else libelle, bloc=bloc, groupe=groupe,
                categorie=str(b) if b is not None else None,
                prevu=None if imprevus else lire_prevu(c), montants=lire_montants(d),
                commentaire=" ".join(str(f).split()) if f not in (None, "") else None)


# ---------------------------------------------------------------- accès à l'appli (Supabase)
class HorsLigne(Exception):
    pass


class ErreurAppli(Exception):
    pass


class Appli:
    def __init__(self, config, enregistrer_config):
        self.config, self.enregistrer_config = config, enregistrer_config
        self.jeton, self.expire = None, 0

    def _http(self, methode, chemin, corps=None, entetes=None, auth=True):
        h = {"apikey": SUPABASE_KEY, "Accept": "application/json"}
        if corps is not None:
            h["Content-Type"] = "application/json"
        if auth:
            h["Authorization"] = f"Bearer {self._jeton()}"
        h.update(entetes or {})
        req = urllib.request.Request(SUPABASE_URL + chemin, method=methode, headers=h,
                                     data=None if corps is None else json.dumps(corps).encode())
        try:
            with urllib.request.urlopen(req, timeout=30) as rep:
                texte = rep.read().decode()
                return json.loads(texte) if texte else None
        except urllib.error.HTTPError as e:
            detail = e.read().decode(errors="replace")[:300]
            raise ErreurAppli(f"{methode} {chemin.split('?')[0]} : {e.code} {detail}") from None
        except (urllib.error.URLError, TimeoutError, ConnectionError, OSError) as e:
            raise HorsLigne(str(e)) from None

    def connexion(self, email, mot_de_passe):
        j = self._http("POST", "/auth/v1/token?grant_type=password",
                       {"email": email, "password": mot_de_passe}, auth=False)
        self._garder(j)

    def _garder(self, j):
        self.jeton, self.expire = j["access_token"], time.time() + int(j.get("expires_in", 3600)) - 60
        self.config["refresh_token"] = j["refresh_token"]
        self.config["email"] = (j.get("user") or {}).get("email", self.config.get("email"))
        self.enregistrer_config(self.config)

    def _jeton(self):
        if self.jeton and time.time() < self.expire:
            return self.jeton
        if not self.config.get("refresh_token"):
            raise ErreurAppli("Pas connecté : lancez « budget_sync.py connexion »")
        j = self._http("POST", "/auth/v1/token?grant_type=refresh_token",
                       {"refresh_token": self.config["refresh_token"]}, auth=False)
        self._garder(j)
        return self.jeton

    def get(self, chemin):
        return self._http("GET", "/rest/v1/" + chemin)

    def post(self, chemin, corps, retour=True):
        return self._http("POST", "/rest/v1/" + chemin, corps,
                          {"Prefer": "return=representation" if retour else "return=minimal"})

    def patch(self, chemin, corps):
        return self._http("PATCH", "/rest/v1/" + chemin, corps, {"Prefer": "return=minimal"})

    def delete(self, chemin):
        return self._http("DELETE", "/rest/v1/" + chemin, None, {"Prefer": "return=minimal"})

    def versions(self):
        return {v["mois_id"]: v["version"] for v in self.get("versions_mois?select=mois_id,version")}

    def charger(self, ids=None):
        mois = {m["id"]: m for m in self.get("mois?select=*&order=id.asc")}
        ids = [i for i in (ids or mois) if i in mois]
        lignes = []
        if ids:
            lignes = self.get(f"lignes?select=*,operations(*)&mois_id=in.({','.join(ids)})&order=ordre.asc")
        par_mois = {i: {"mois": mois[i], "lignes": {}} for i in ids}
        for l in lignes:
            cle = "imprevus" if l.get("code") == "imprevus" else normaliser(l["libelle"])
            l["operations"] = sorted(l.get("operations") or [], key=lambda o: (o.get("created_at") or "", o["id"]))
            par_mois[l["mois_id"]]["lignes"][cle] = l
        return mois, par_mois


# ---------------------------------------------------------------- plan d'actions
class Plan:
    def __init__(self):
        self.appli = []      # (description, fonction)
        self.excel = []      # descriptions
        self.excel_modifie = False
        self.messages = []

    def a(self, description, fonction):
        self.appli.append((description, fonction))

    def x(self, description):
        self.excel.append(description)
        self.excel_modifie = True


def resume_montants(cs):
    return ", ".join(f"{en_euros(c):.2f} €".replace(".", ",") for c in cs)


# ---------------------------------------------------------------- synchronisation
class Synchro:
    def __init__(self, config, essai=False):
        self.config, self.essai = config, essai
        self.fichier = Path(config["fichier"]).expanduser()
        self.appli = Appli(config, enregistrer_config)
        self.fichier_etat = DOSSIER_DONNEES / "etat.json"
        self.etat = json.loads(self.fichier_etat.read_text()) if self.fichier_etat.exists() else {"mois": {}}
        self.attente_excel = False

    # ---- fichier
    def fichier_ouvert(self):
        verrou = self.fichier.parent / f".~lock.{self.fichier.name}#"
        return verrou.exists()

    def sauvegarder_classeur(self):
        dossier = DOSSIER_DONNEES / "sauvegardes"
        dossier.mkdir(parents=True, exist_ok=True)
        shutil.copy2(self.fichier, dossier / f"{self.fichier.stem}_{datetime.now():%Y-%m-%d_%H-%M-%S}.xlsx")
        anciennes = sorted(dossier.glob("*.xlsx"))
        for f in anciennes[:-SAUVEGARDES_GARDEES]:
            f.unlink()

    def ecrire_classeur(self, wb, mtime_lu):
        if os.path.getmtime(self.fichier) != mtime_lu:
            return False  # modifié pendant la synchro : on recommencera
        self.sauvegarder_classeur()
        wb.calculation = CalcProperties(fullCalcOnLoad=True)
        fd, tmp = tempfile.mkstemp(suffix=".xlsx", dir=self.fichier.parent, prefix=".budget-sync-")
        os.close(fd)
        try:
            wb.save(tmp)
            shutil.copymode(self.fichier, tmp)
            os.replace(tmp, self.fichier)
        finally:
            if os.path.exists(tmp):
                os.unlink(tmp)
        return True

    # ---- un cycle complet
    def cycle(self):
        mtime = os.path.getmtime(self.fichier)
        wb = openpyxl.load_workbook(self.fichier)
        feuilles = {}
        for ws in wb.worksheets:
            mid = mois_de_feuille(ws.title)
            if mid:
                feuilles[mid] = Feuille(ws, mid)
        versions = self.appli.versions()
        mois_db, donnees = self.appli.charger()
        plan = Plan()
        nouveaux_etats = {}
        ouvert = self.fichier_ouvert()

        tous = sorted(set(feuilles) | set(mois_db))
        premiere_feuille = min(feuilles) if feuilles else None
        for mid in tous:
            f, d = feuilles.get(mid), donnees.get(mid)
            base = self.etat["mois"].get(mid)
            if f and d:
                nouveaux_etats[mid] = self.fusionner_mois(f, d, base, plan)
            elif f and not d:
                nouveaux_etats[mid] = self.creer_mois_appli(f, mois_db, donnees, plan)
            elif d and not f and premiere_feuille and mid > premiere_feuille:
                nouveaux_etats[mid] = self.creer_feuille(wb, feuilles, mid, d, plan)

        # 1) l'appli (toujours possible)
        for description, fonction in plan.appli:
            log.info("Appli : %s", description)
            if not self.essai:
                fonction()
        # 2) le classeur (seulement s'il n'est pas ouvert)
        excel_ecrit = False
        if plan.excel_modifie:
            for description in plan.excel:
                log.info("Excel : %s", description)
            if self.essai:
                pass
            elif ouvert:
                if not self.attente_excel:
                    notifier("Des changements de l'appli attendent la fermeture du fichier Excel.")
                self.attente_excel = True
            else:
                excel_ecrit = self.ecrire_classeur(wb, mtime)
                self.attente_excel = not excel_ecrit
                if excel_ecrit:
                    log.info("Classeur mis à jour (%d changement(s))", len(plan.excel))
        for m in plan.messages:
            notifier(m)
        if self.essai:
            return plan
        # 3) mémoriser l'état de référence pour la prochaine fois
        for mid, (etat_si_ecrit, etat_sinon) in nouveaux_etats.items():
            if etat_si_ecrit is None:
                continue
            ecrit = excel_ecrit or not plan.excel_modifie
            self.etat["mois"][mid] = etat_si_ecrit if ecrit else etat_sinon
        if plan.appli:
            # relire les identifiants des opérations créées
            _, donnees = self.appli.charger([m for m in nouveaux_etats])
            for mid, d in donnees.items():
                for cle, l in d["lignes"].items():
                    e = self.etat["mois"].get(mid, {}).get("lignes", {}).get(cle)
                    if e is not None:
                        e["ids"] = [o["id"] for o in l["operations"]]
        self.etat["versions"] = self.appli.versions()
        self.etat["fichier_mtime"] = os.path.getmtime(self.fichier)
        self.fichier_etat.parent.mkdir(parents=True, exist_ok=True)
        tmp = self.fichier_etat.with_suffix(".tmp")
        tmp.write_text(json.dumps(self.etat, ensure_ascii=False, indent=1))
        os.replace(tmp, self.fichier_etat)
        return plan

    # ---- fusion d'un mois présent des deux côtés
    def fusionner_mois(self, f, d, base, plan):
        ws, mid = f.ws, f.mois_id
        bl = (base or {}).get("lignes", {})
        premiere = base is None
        etat_ecrit = {"lignes": {}}
        etat_excel = {"lignes": {}}
        dl = dict(d["lignes"])
        el = dict(f.lignes)

        # renommage dans Excel : même rang, nom différent
        for cle_e in [k for k in el if k not in dl]:
            rang = el[cle_e]["rang"]
            for cle_d, l in list(dl.items()):
                if cle_d not in el and l["ordre"] == rang and cle_d != "imprevus":
                    nouveau = el[cle_e]["libelle"]
                    plan.a(f"{nom_mois(mid)} : « {l['libelle']} » renommée « {nouveau} »",
                           lambda i=l["id"], n=nouveau: self.appli.patch(f"lignes?id=eq.{i}", {"libelle": n}))
                    dl[cle_e] = dict(l, libelle=nouveau)
                    del dl[cle_d]
                    if cle_d in bl:
                        bl = dict(bl)
                        bl[cle_e] = bl.pop(cle_d)
                    break

        for cle in sorted(set(el) | set(dl)):
            e, l, b = el.get(cle), dl.get(cle), bl.get(cle)
            if e and l:
                ecrit, sinon = self.fusionner_ligne(ws, mid, e, l, b, premiere, plan)
                etat_ecrit["lignes"][cle], etat_excel["lignes"][cle] = ecrit, sinon
            elif e and not l:
                self.creer_ligne_appli(mid, e, plan)
                snap = {"prevu": e["prevu"], "montants": e["montants"] or [], "commentaire": e["commentaire"], "ids": []}
                etat_ecrit["lignes"][cle] = etat_excel["lignes"][cle] = snap
            elif l and not e:
                ids = [o["id"] for o in l["operations"]]
                if b is not None and set(ids) <= set(b.get("ids", [])):
                    plan.a(f"{nom_mois(mid)} : ligne « {l['libelle']} » supprimée (retirée du fichier Excel)",
                           lambda i=l["id"]: self.appli.delete(f"lignes?id=eq.{i}"))
                elif b is not None:
                    plan.messages.append(f"« {l['libelle']} » ({nom_mois(mid)}) a été retirée du fichier Excel mais contient "
                                         "des dépenses saisies sur téléphone : remettez la ligne dans Excel.")
                    etat_ecrit["lignes"][cle] = etat_excel["lignes"][cle] = b
                else:
                    plan.messages.append(f"La ligne « {l['libelle']} » ({nom_mois(mid)}) existe dans l'appli mais pas dans "
                                         "Excel : ajoutez-la dans la feuille avec le même nom.")

        # valeurs du mois : provision imprévus et report
        mois = d["mois"]
        tB = (base or {}).get("taux", float(mois["taux_imprevus"]))
        tD = float(mois["taux_imprevus"])
        tE = f.taux
        taux = tD
        if tE is not None and f.rang_imprevus:
            taux = fusion_valeur(tB, tE, tD) if not premiere else tE
            if abs(taux - tD) > 1e-9:
                plan.a(f"{nom_mois(mid)} : provision imprévus {taux * 100:g} %",
                       lambda t=taux: self.appli.patch(f"mois?id=eq.{mid}", {"taux_imprevus": t}))
            if abs(taux - tE) > 1e-9 and f.rang_total:
                ws.cell(f.rang_imprevus, 3).value = f"=ROUND(C{f.rang_total}*{taux * 100:g}%,2)" if taux else None
                ws.cell(f.rang_imprevus, 1).value = f"IMPRÉVUS {str(round(taux * 100, 2)).replace('.', ',').rstrip('0').rstrip(',')}%"
                plan.x(f"{nom_mois(mid)} : provision imprévus {taux * 100:g} %")
        rB = (base or {}).get("report", float(mois["report_solde"]))
        rD = float(mois["report_solde"])
        rE = f.report
        report = rD
        if rE is not None:
            report = fusion_valeur(rB, rE, rD) if not premiere else rE
            if abs(report - rD) > 0.004:
                plan.a(f"{nom_mois(mid)} : solde reporté {report:+.2f} €",
                       lambda r=report: self.appli.patch(f"mois?id=eq.{mid}", {"report_solde": r}))
        if (rE is None and abs(rD) > 0.004) or (rE is not None and abs(report - rE) > 0.004):
            self.ecrire_report(f, report, mois.get("note_report"))
            plan.x(f"{nom_mois(mid)} : solde reporté {report:+.2f} €")
        for etat in (etat_ecrit, etat_excel):
            etat["taux"], etat["report"] = taux, report
        if rE is not None:
            etat_excel["report"] = rE
        if tE is not None:
            etat_excel["taux"] = tE
        return etat_ecrit, etat_excel

    def ecrire_report(self, f, report, note):
        ws = f.ws
        if f.rang_solde is None:
            return
        rang = f.rang_report or (f.rang_solde + 2)
        signe = "-" if report < 0 else "+"
        ws.cell(rang, 1).value = LIBELLE_REPORT
        ws.cell(rang, 4).value = f"=D{f.rang_solde}{signe}{texte_montant(centimes(report))}"
        ws.cell(rang, 4).number_format = ws.cell(f.rang_solde, 4).number_format
        if note:
            ws.cell(rang, 6).value = note
        f.rang_report = rang

    def fusionner_ligne(self, ws, mid, e, l, b, premiere, plan):
        nom = f"{nom_mois(mid)}, {l['libelle']}"
        ops = l["operations"]
        D = Counter(centimes(o["montant"]) for o in ops)
        ids_base = set((b or {}).get("ids", []))
        excel_montants = e["montants"]
        resultat_montants = None
        if excel_montants is None:
            plan.messages.append(f"{nom} : le « coût réel » n'est pas une simple somme, ligne non synchronisée.")
            m_ecrit = m_excel = [centimes(o["montant"]) for o in ops]
        else:
            E = Counter(excel_montants)
            B = Counter(b["montants"]) if b else (E & D)
            supE, ajE, supD, ajD = B - E, E - B, B - D, D - B
            M = (B - (supE | supD)) + ajE + ajD
            # côté appli
            a_ajouter = M - D
            a_retirer = D - M
            if a_ajouter:
                lignes_ops = [{"ligne_id": l["id"], "montant": en_euros(c), "date": date.today().isoformat()}
                              for c in a_ajouter.elements()]
                plan.a(f"{nom} : + {resume_montants(list(a_ajouter.elements()))}",
                       lambda corps=lignes_ops: self.appli.post("operations", corps, retour=False))
            if a_retirer:
                choix, reste = [], Counter(a_retirer)
                for o in sorted(ops, key=lambda o: (o["id"] not in ids_base, o.get("created_at") or "")):
                    c = centimes(o["montant"])
                    if reste[c] > 0:
                        choix.append(o["id"])
                        reste[c] -= 1
                plan.a(f"{nom} : − {resume_montants(list(a_retirer.elements()))}",
                       lambda ids=choix: self.appli.delete(f"operations?id=in.({','.join(ids)})"))
            # côté Excel : on garde l'ordre de la formule, on retire puis on ajoute à la fin
            liste = list(excel_montants)
            for c, n in (E - M).items():
                for _ in range(n):
                    i = len(liste) - 1 - liste[::-1].index(c)
                    liste.pop(i)
            for o in ops:
                c = centimes(o["montant"])
                if (M - Counter(liste))[c] > 0:
                    liste.append(c)
            for c in (M - Counter(liste)).elements():
                liste.append(c)
            if liste != excel_montants:
                ws.cell(e["rang"], 4).value = valeur_montants(liste)
                plan.x(f"{nom} : coût réel {valeur_montants(liste)}")
            m_ecrit, m_excel = liste, excel_montants

        # montant prévu
        p_ecrit = p_excel = e["prevu"]
        if l.get("code") != "imprevus" and e["prevu"] != "?":
            pD = float(l["prevu"]) if l["prevu"] is not None else None
            pE = e["prevu"]
            pB = (b or {}).get("prevu", pD) if not premiere else pD
            p = fusion_valeur(pB, pE, pD)
            if p != pD:
                plan.a(f"{nom} : prévu {p} €", lambda v=p: self.appli.patch(f"lignes?id=eq.{l['id']}", {"prevu": v}))
            if p != pE:
                ws.cell(e["rang"], 3).value = p
                plan.x(f"{nom} : prévu {p} €")
            p_ecrit = p

        # commentaire (+ notes des achats saisis sur téléphone)
        cD, cE = l.get("commentaire") or None, e["commentaire"]
        cB = (b or {}).get("commentaire", cD) if not premiere else cD
        c = fusion_valeur(cB, cE, cD)
        if not premiere:
            nouvelles = [o["note"].strip() for o in ops if o["id"] not in ids_base and (o.get("note") or "").strip()]
            for n in nouvelles:
                if not c or n.casefold() not in c.casefold():
                    c = f"{c} + {n}" if c else n
        if c != cD:
            plan.a(f"{nom} : commentaire", lambda v=c: self.appli.patch(f"lignes?id=eq.{l['id']}", {"commentaire": v}))
        if c != cE:
            ws.cell(e["rang"], 6).value = c
            plan.x(f"{nom} : commentaire « {c or ''} »")

        ecrit = {"prevu": p_ecrit, "montants": m_ecrit, "commentaire": c, "ids": [o["id"] for o in ops]}
        sinon = {"prevu": e["prevu"], "montants": m_excel, "commentaire": cE, "ids": [o["id"] for o in ops]}
        return ecrit, sinon

    def creer_ligne_appli(self, mid, e, plan):
        def faire():
            [ligne] = self.appli.post("lignes", {
                "mois_id": mid, "bloc": e["bloc"], "groupe": e["groupe"], "libelle": e["libelle"],
                "categorie": e["categorie"], "prevu": e["prevu"] if e["prevu"] != "?" else None,
                "code": "imprevus" if e["libelle"] == "Imprévus" else None,
                "commentaire": e["commentaire"], "ordre": e["rang"]})
            if e["montants"]:
                self.appli.post("operations", [{"ligne_id": ligne["id"], "montant": en_euros(c)} for c in e["montants"]], retour=False)
        plan.a(f"{nom_mois(mid)} : nouvelle ligne « {e['libelle']} »", faire)

    # ---- nouvelle feuille dans Excel -> nouveau mois dans l'appli
    def creer_mois_appli(self, f, mois_db, donnees, plan):
        mid = f.mois_id
        report = f.report
        if report is None or (f.rang_report is None):
            precedent = max((m for m in donnees if m < mid), default=None)
            report = 0.0
            if precedent:
                report = solde_avec_report(donnees[precedent])
            if f.rang_solde:
                self.ecrire_report(f, report, f"Solde fin {nom_mois(precedent)}" if precedent else None)
                plan.x(f"{nom_mois(mid)} : solde reporté {report:+.2f} €")
        taux = f.taux if f.taux is not None else 0.035

        def faire():
            self.appli.post("mois", {"id": mid, "report_solde": report, "taux_imprevus": taux}, retour=False)
            corps = [{"mois_id": mid, "bloc": e["bloc"], "groupe": e["groupe"], "libelle": e["libelle"],
                      "categorie": e["categorie"], "prevu": e["prevu"] if e["prevu"] != "?" else None,
                      "code": "imprevus" if cle == "imprevus" else None, "commentaire": e["commentaire"],
                      "ordre": e["rang"]} for cle, e in f.lignes.items()]
            lignes = self.appli.post("lignes", corps)
            ops = []
            for l in lignes:
                cle = "imprevus" if l.get("code") == "imprevus" else normaliser(l["libelle"])
                ops += [{"ligne_id": l["id"], "montant": en_euros(c)} for c in (f.lignes[cle]["montants"] or [])]
            if ops:
                self.appli.post("operations", ops, retour=False)
        plan.a(f"nouveau mois {nom_mois(mid)} créé depuis la feuille Excel ({len(f.lignes)} lignes)", faire)
        snap = {"taux": taux, "report": report, "lignes": {
            cle: {"prevu": e["prevu"], "montants": e["montants"] or [], "commentaire": e["commentaire"], "ids": []}
            for cle, e in f.lignes.items()}}
        return snap, snap

    # ---- nouveau mois dans l'appli -> nouvelle feuille dans Excel
    def creer_feuille(self, wb, feuilles, mid, d, plan):
        precedents = [m for m in feuilles if m < mid]
        modele = feuilles[max(precedents)].ws if precedents else next(
            (ws for ws in wb.worksheets if "vierge" in ws.title.casefold()), None)
        if modele is None:
            plan.messages.append(f"Impossible de créer la feuille de {nom_mois(mid)} : aucun modèle.")
            return None, None
        ws = wb.copy_worksheet(modele)
        ws.title = nom_feuille(mid)
        wb.move_sheet(ws, offset=-(len(wb.worksheets) - 1 - wb.worksheets.index(modele) - 1))
        for plage in modele.conditional_formatting:
            for regle in plage.rules:
                ws.conditional_formatting.add(str(plage.sqref), regle)
        ws.sheet_view.tabSelected = False
        f = Feuille(ws, mid)
        par_rang = {e["rang"]: (cle, e) for cle, e in f.lignes.items()}
        for e in f.lignes.values():                       # on vide le réel et les commentaires du modèle
            ws.cell(e["rang"], 4).value = None
            ws.cell(e["rang"], 6).value = None
        mois = d["mois"]
        lignes_snap = {}
        for cle, l in d["lignes"].items():
            rang = l["ordre"]
            if rang not in par_rang:
                plan.messages.append(f"{nom_mois(mid)} : la ligne « {l['libelle']} » n'a pas de place dans la feuille, ajoutez-la à la main.")
                continue
            if l.get("code") != "imprevus":
                ws.cell(rang, 1).value = l["libelle"]
                ws.cell(rang, 3).value = float(l["prevu"]) if l["prevu"] is not None else None
            if l.get("categorie"):
                ws.cell(rang, 2).value = l["categorie"]
            montants = [centimes(o["montant"]) for o in l["operations"]]
            ws.cell(rang, 4).value = valeur_montants(montants)
            com = l.get("commentaire")
            notes = [o["note"] for o in l["operations"] if (o.get("note") or "").strip()]
            com = " + ".join([x for x in [com] + notes if x]) or None
            ws.cell(rang, 6).value = com
            lignes_snap[cle] = {"prevu": float(l["prevu"]) if l["prevu"] is not None else None, "montants": montants,
                                "commentaire": com, "ids": [o["id"] for o in l["operations"]]}
        taux = float(mois["taux_imprevus"])
        if f.rang_imprevus and f.rang_total:
            ws.cell(f.rang_imprevus, 3).value = f"=ROUND(C{f.rang_total}*{taux * 100:g}%,2)" if taux else None
        report = float(mois["report_solde"])
        self.ecrire_report(f, report, mois.get("note_report"))
        plan.x(f"nouvelle feuille « {ws.title} » ({len(lignes_snap)} lignes)")
        snap = {"taux": taux, "report": report, "lignes": lignes_snap}
        return snap, None

    # ---- boucle
    def surveiller(self):
        dernier_mtime, dernieres_versions, dernier_regard_appli = None, None, 0
        derniere_erreur = None
        log.info("Surveillance de %s", self.fichier)
        while True:
            erreur = None
            try:
                mtime = os.path.getmtime(self.fichier)
                a_faire = mtime != dernier_mtime or self.attente_excel and not self.fichier_ouvert()
                if time.time() - dernier_regard_appli >= INTERVALLE_APPLI:
                    dernier_regard_appli = time.time()
                    v = self.appli.versions()
                    if v != dernieres_versions:
                        a_faire = True
                if a_faire:
                    self.cycle()
                    dernier_mtime = os.path.getmtime(self.fichier)
                    dernieres_versions = self.appli.versions()
            except HorsLigne as e:
                erreur = f"Pas de réseau ({e}), nouvel essai plus tard"
            except FileNotFoundError:
                erreur = f"Fichier introuvable : {self.fichier}"
            except ErreurAppli as e:
                erreur = f"Erreur avec l'appli : {e}"
                time.sleep(60)
            except Exception as e:  # on ne s'arrête jamais pour une erreur ponctuelle
                log.exception("Erreur inattendue : %s", e)
                time.sleep(30)
            if erreur and erreur != derniere_erreur:
                (notifier if erreur.startswith("Erreur") else log.warning)(erreur)
            elif not erreur and derniere_erreur:
                log.info("Tout refonctionne.")
            derniere_erreur = erreur
            time.sleep(INTERVALLE_FICHIER)


def solde_avec_report(d):
    lignes = d["lignes"].values()
    base = sum(float(l["prevu"] or 0) for l in lignes if l["bloc"] in ("fixe", "variable"))
    reel = lambda l: sum(float(o["montant"]) for o in l["operations"])
    dep = sum(reel(l) for l in lignes if l["bloc"] != "revenu")
    rev = sum(reel(l) for l in lignes if l["bloc"] == "revenu")
    return round(rev - dep + float(d["mois"]["report_solde"] or 0), 2)


# ---------------------------------------------------------------- configuration
def lire_config():
    f = DOSSIER_CONFIG / "config.json"
    return json.loads(f.read_text()) if f.exists() else {}


def enregistrer_config(config):
    DOSSIER_CONFIG.mkdir(parents=True, exist_ok=True)
    f = DOSSIER_CONFIG / "config.json"
    tmp = f.with_suffix(".tmp")
    tmp.write_text(json.dumps(config, indent=1))
    os.chmod(tmp, 0o600)
    os.replace(tmp, f)


def commande_connexion(args):
    config = lire_config()
    print("Connexion de ce PC à l'appli Budget.")
    fichier = input(f"Chemin du fichier Excel [{config.get('fichier', FICHIER_DEFAUT)}] : ").strip()
    config["fichier"] = str(Path(fichier or config.get("fichier", FICHIER_DEFAUT)).expanduser())
    if not Path(config["fichier"]).exists():
        sys.exit(f"Fichier introuvable : {config['fichier']}")
    email = input(f"E-mail [{config.get('email', '')}] : ").strip() or config.get("email", "")
    mdp = getpass.getpass("Mot de passe (il n'est pas enregistré) : ")
    appli = Appli(config, enregistrer_config)
    try:
        appli.connexion(email, mdp)
    except ErreurAppli as e:
        sys.exit(f"Connexion refusée : {e}")
    if not appli.get("membres?select=email"):
        sys.exit("Ce compte n'a pas accès au budget du foyer.")
    print("Connecté. Ce PC peut maintenant synchroniser le fichier.")


def main():
    parser = argparse.ArgumentParser(description="Synchronise le classeur Excel du budget avec l'appli.")
    parser.add_argument("commande", nargs="?", default="surveiller", choices=["surveiller", "connexion", "essai", "une-fois"])
    parser.add_argument("-v", "--bavard", action="store_true")
    args = parser.parse_args()
    DOSSIER_DONNEES.mkdir(parents=True, exist_ok=True)
    logging.basicConfig(level=logging.DEBUG if args.bavard else logging.INFO,
                        format="%(asctime)s %(levelname)s %(message)s",
                        handlers=[logging.StreamHandler(sys.stdout),
                                  logging.FileHandler(DOSSIER_DONNEES / "journal.log", encoding="utf-8")])
    if args.commande == "connexion":
        return commande_connexion(args)
    config = lire_config()
    if not config.get("refresh_token") or not config.get("fichier"):
        sys.exit("Ce PC n'est pas encore connecté : lancez « python3 budget_sync.py connexion ».")
    s = Synchro(config, essai=args.commande == "essai")
    if args.commande in ("essai", "une-fois"):
        plan = s.cycle()
        if args.commande == "essai":
            print(f"\nEssai : {len(plan.appli)} changement(s) côté appli, {len(plan.excel)} côté Excel. Rien n'a été modifié.")
        for m in plan.messages:
            print("Attention :", m)
        return None
    return s.surveiller()


if __name__ == "__main__":
    main()
